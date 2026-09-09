#!/usr/bin/env -S uv run --script
# /// script
# requires-python = "==3.12.*"
# dependencies = [
#   "markitdown[pdf,docx,pptx,xlsx,xls]==0.1.6",
#   "openpyxl==3.1.5",
# ]
# ///
"""Convert a binary document to LLM-readable Markdown.

    uv run --script --frozen scripts/convert_to_md.py <input> [--out <output.md>]

Handles the six Office formats via MarkItDown, plus the intake formats the
harness accepts directly: plain text, CSV/TSV, JSON, .eml, and WebVTT/SRT
transcripts. Audio and images are refused with a clear explanation rather than
a confusing failure — they need a transcription or description service first.

For .xlsx/.xlsm, --formulas extracts literal cell formulas (=B2*C2) instead of
MarkItDown's computed-value tables. Use it when you need a workbook's *logic*
rather than its numbers -- reverse-engineering a client's spreadsheet model into
an ontology, for instance. Run both passes to separate files when you need both.

Writes to <name>.md in the current directory unless --out is given. Refuses to
write into datasources/, which the harness treats as read-only.

Prints the written path to stdout; diagnostics to stderr.
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

SUPPORTED = frozenset({".pdf", ".docx", ".pptx", ".xlsx", ".xlsm", ".xls"})
FORMULA_CAPABLE = frozenset({".xlsx", ".xlsm"})
READ_ONLY_DIR = "datasources"

# Formats intake accepts without MarkItDown. A transcript is roughly half
# timing metadata, so stripping it before an extraction pass is worth doing.
TRANSCRIPT = frozenset({".vtt", ".srt"})
PLAIN = frozenset({".txt", ".md", ".csv", ".tsv", ".log", ".json", ".jsonl"})
EMAIL = frozenset({".eml"})

# Formats that need a service the harness does not have. Saying so plainly
# beats a confusing failure three steps later.
NEEDS_SERVICE = {
    ".m4a": "audio — needs transcription",
    ".mp3": "audio — needs transcription",
    ".wav": "audio — needs transcription",
    ".ogg": "audio — needs transcription",
    ".mp4": "video — needs transcription",
    ".mov": "video — needs transcription",
    ".png": "image — needs a description pass",
    ".jpg": "image — needs a description pass",
    ".jpeg": "image — needs a description pass",
    ".heic": "image — needs a description pass",
    ".webp": "image — needs a description pass",
    ".gif": "image — needs a description pass",
}


def fail(msg: str) -> None:
    print(f"error: {msg}", file=sys.stderr)
    raise SystemExit(1)


def check_output_path(out: Path) -> None:
    """Refuse to write anywhere under a datasources/ directory."""
    parts = {p.lower() for p in out.resolve().parts}
    if READ_ONLY_DIR in parts:
        fail(
            f"refusing to write into {READ_ONLY_DIR}/ -- it is read-only to the "
            "harness. Write into the engagement folder instead."
        )


def convert_transcript(src: Path) -> str:
    """Strip WebVTT/SRT to plain speech.

    Cue numbers, timecodes and inline tags are noise to an extraction pass and
    a surprising share of the tokens.
    """
    out: list[str] = []
    last = ""
    for raw in src.read_text(encoding="utf-8", errors="replace").splitlines():
        s = raw.strip()
        if not s or s == "WEBVTT" or s.startswith("NOTE "):
            continue
        if s.isdigit() or "-->" in s:
            continue
        clean = re.sub(r"<[^>]+>", "", s).strip()
        if not clean or clean == last:
            continue
        last = clean
        out.append(clean)
    lines = [f"# Transcript — {src.name}", ""]
    lines.extend(out)
    return "\n".join(lines) + "\n"


def convert_email(src: Path) -> str:
    """Headers worth keeping, then the plain-text body."""
    from email import policy
    from email.parser import BytesParser

    msg = BytesParser(policy=policy.default).parsebytes(src.read_bytes())
    head = [f"# Email — {msg.get('subject', '(no subject)')}", ""]
    for k in ("From", "To", "Cc", "Date"):
        v = msg.get(k)
        if v:
            head.append(f"- **{k}:** {v}")
    head.append("")
    body = msg.get_body(preferencelist=("plain", "html"))
    text = body.get_content() if body else ""
    if body is not None and body.get_content_type() == "text/html":
        text = re.sub(r"<[^>]+>", " ", text)
        text = re.sub(r"\s+", " ", text)
    return "\n".join(head) + text.strip() + "\n"


def convert_plain(src: Path) -> str:
    """Passthrough, so every intake path produces a .md the same way."""
    return src.read_text(encoding="utf-8", errors="replace")


def convert_values(src: Path) -> str:
    from markitdown import MarkItDown

    result = MarkItDown().convert(str(src))
    return result.text_content


def convert_formulas(src: Path) -> str:
    """Dump literal cell formulas per sheet, with coordinates.

    Cells holding only a static value are skipped -- the point of this mode is
    the logic. A workbook with no formulas produces an explicit note rather than
    an empty file, so the caller can tell "no formulas" from "conversion failed".
    """
    from openpyxl import load_workbook

    wb = load_workbook(filename=str(src), data_only=False, read_only=True)
    out: list[str] = [f"# Cell formulas — {src.name}", ""]
    found = 0

    for ws in wb.worksheets:
        rows: list[str] = []
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("="):
                    rows.append(f"| `{cell.coordinate}` | `{value}` |")
                    found += 1
        if rows:
            out.append(f"## Sheet: {ws.title}")
            out.append("")
            out.append("| Cell | Formula |")
            out.append("|---|---|")
            out.extend(rows)
            out.append("")

    wb.close()

    if not found:
        out.append(
            "_No formulas found in this workbook. Every cell holds a static "
            "value; re-run without `--formulas` to extract those._"
        )
    else:
        out.insert(2, f"_{found} formula cells across {len(wb.sheetnames)} sheets._")
        out.insert(3, "")

    return "\n".join(out)


def main() -> None:
    ap = argparse.ArgumentParser(
        description="Convert a binary document to Markdown for the FDE harness."
    )
    ap.add_argument("input", type=Path, help="PDF, DOCX, PPTX, XLSX, XLSM or XLS")
    ap.add_argument("--out", type=Path, help="output path (default: <name>.md here)")
    ap.add_argument(
        "--formulas",
        action="store_true",
        help="xlsx/xlsm only: extract literal cell formulas instead of values",
    )
    args = ap.parse_args()

    src: Path = args.input
    if not src.is_file():
        fail(f"not a file: {src}")

    ext = src.suffix.lower()
    if ext in NEEDS_SERVICE:
        fail(
            f"{ext} is {NEEDS_SERVICE[ext]}. The harness has no transcription or "
            "vision service. Produce a text version first and drop that in the "
            "same evidence folder — the class is set by the folder, so the "
            "transcript keeps the standing of the original."
        )
    known = SUPPORTED | TRANSCRIPT | PLAIN | EMAIL
    if ext not in known:
        fail(f"unsupported extension {ext!r}. Supported: {', '.join(sorted(known))}")

    if args.formulas and ext not in FORMULA_CAPABLE:
        fail(
            f"--formulas needs .xlsx or .xlsm, got {ext!r}. Other formats carry no "
            "formulas; run without --formulas."
        )

    out: Path = args.out or Path.cwd() / f"{src.stem}.md"
    check_output_path(out)

    try:
        if args.formulas:
            text = convert_formulas(src)
        elif ext in TRANSCRIPT:
            text = convert_transcript(src)
        elif ext in EMAIL:
            text = convert_email(src)
        elif ext in PLAIN:
            text = convert_plain(src)
        else:
            text = convert_values(src)
    except Exception as exc:  # noqa: BLE001 - surface the real cause to the operator
        fail(f"conversion failed: {type(exc).__name__}: {exc}")

    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(text, encoding="utf-8")
    print(out)


if __name__ == "__main__":
    main()
